from flask import Flask, render_template, request, redirect, session, jsonify, send_file
from io import BytesIO
from pathlib import Path
import os
import json
import re
import subprocess
import sys
import tempfile
import time

from database import initialize_database, save_student, save_result, get_leaderboard, get_result, delete_result, clear_results
from questions import SETS


app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "cyber-debug-secret-key-change-this"
)


# ============================================================
# SETTINGS
# ============================================================

TEST_DURATION = 25 * 60

ADMIN_PASSWORD = os.environ.get(
    "ADMIN_PASSWORD",
    "ADITYA"
)

DUSTBIN_PASSWORD = os.environ.get(
    "DUSTBIN_PASSWORD",
    "HEMA"
)

SHOW_ANSWERS_PASSWORD = os.environ.get(
    "SHOW_ANSWERS_PASSWORD",
    "show@123"
)

MAX_SCORE_PER_QUESTION = 10

# Each debugging question has exactly four hidden tests.
# Each passed hidden test is worth 2.5 marks.
POINTS_PER_TEST = 2.5

TESTS_PER_QUESTION = 4


# ============================================================
# TIMER
# ============================================================

def current_elapsed():
    start = session.get("test_start_time")

    if not start:
        return 0

    return max(
        0,
        int(time.time() - float(start))
    )


def admin_required():
    return session.get("admin_authenticated") is True


# ============================================================
# EVALUATE QUESTION
# ============================================================

def evaluate_question(question, submitted_code):

    if not submitted_code.strip():
        return (
            0,
            0,
            "No answer submitted"
        )

    passed = 0
    temp_path = None

    try:

        # ------------------------------------------------------
        # Create temporary Python file
        # ------------------------------------------------------

        with tempfile.NamedTemporaryFile(
            mode="w",
            suffix=".py",
            delete=False,
            encoding="utf-8"
        ) as f:

            f.write(submitted_code)
            temp_path = f.name


        # ------------------------------------------------------
        # Run all hidden tests
        # ------------------------------------------------------

        for test in question["tests"]:

            try:

                # Your questions.py uses:
                # ("input", "expected output")

                if isinstance(test, (tuple, list)):

                    test_input = test[0]
                    expected_output = test[1]

                # Also support dictionary format

                elif isinstance(test, dict):

                    test_input = test.get(
                        "input",
                        ""
                    )

                    expected_output = test.get(
                        "output",
                        ""
                    )

                else:

                    continue


                # ------------------------------------------------
                # Execute participant code
                # ------------------------------------------------

                completed = subprocess.run(

                    [
                        sys.executable,
                        temp_path
                    ],

                    input=test_input,

                    text=True,

                    capture_output=True,

                    timeout=2,

                    cwd=str(
                        Path(temp_path).parent
                    ),

                    env=os.environ.copy()

                )


                # ------------------------------------------------
                # Compare output
                # ------------------------------------------------

                actual = completed.stdout.strip()

                expected = str(
                    expected_output
                ).strip()


                if (
                    completed.returncode == 0
                    and actual == expected
                ):

                    passed += 1


            except subprocess.TimeoutExpired:

                continue

            except (
                OSError,
                ValueError,
                IndexError,
                TypeError
            ):

                continue


        # ------------------------------------------------------
        # Score
        # ------------------------------------------------------

        score = passed * POINTS_PER_TEST

        message = (
            f"{passed}/{TESTS_PER_QUESTION} "
            f"test cases passed"
        )

        return (
            score,
            passed,
            message
        )


    finally:

        # ------------------------------------------------------
        # Delete temporary file
        # ------------------------------------------------------

        if temp_path:

            try:
                os.remove(temp_path)

            except OSError:
                pass


# ============================================================
# LOGIN
# ============================================================

@app.route("/")
def home():

    return render_template("home.html")


@app.route("/login")
def login():

    return render_template("login.html")


# ============================================================
# START
# ============================================================

@app.route("/start", methods=["POST"])
def start():

    name = request.form.get(
        "name",
        ""
    ).strip()

    email = request.form.get(
        "email",
        ""
    ).strip()

    team_name = request.form.get(
        "team_name",
        ""
    ).strip()


    if (
        not name
        or not email
        or not team_name
    ):

        return redirect("/login")


    session.clear()

    session["name"] = name

    session["email"] = email

    session["team_name"] = team_name


    session["student_id"] = save_student(
        name,
        email,
        team_name
    )


    session["stage"] = "instructions"


    return redirect(
        "/instructions"
    )


# ============================================================
# INSTRUCTIONS
# ============================================================

@app.route("/instructions")
def instructions():

    if session.get("stage") != "instructions":

        return redirect("/login")


    return render_template(
        "instructions.html",
        duration=TEST_DURATION
    )


# ============================================================
# SETS
# ============================================================

@app.route("/sets")
def sets_page():

    if session.get("stage") not in (
        "instructions",
        "sets",
        "ready"
    ):

        return redirect("/login")


    session["stage"] = "sets"


    cards = [

        {
            "number": number,
            "name": data["name"],
            "locked": False
        }

        for number, data in SETS.items()

    ]


    return render_template(
        "sets.html",
        sets=cards
    )


# ============================================================
# UNLOCK SET
# ============================================================

@app.route("/unlock", methods=["POST"])
def unlock():

    if session.get("stage") != "sets":

        return redirect("/login")


    try:

        set_number = int(
            request.form.get(
                "set_number",
                "0"
            )
        )

    except ValueError:

        return redirect("/sets")


    password = request.form.get(
        "password",
        ""
    )


    selected = SETS.get(
        set_number
    )


    # --------------------------------------------------------
    # Wrong password
    # --------------------------------------------------------

    if (
        not selected
        or password != selected["password"]
    ):

        cards = [

            {
                "number": number,
                "name": data["name"],
                "locked": False
            }

            for number, data in SETS.items()

        ]


        return render_template(

            "sets.html",

            sets=cards,

            error=(
                "Incorrect password. "
                "Please try again."
            ),

            selected_set=set_number

        )


    # --------------------------------------------------------
    # Correct password
    # --------------------------------------------------------

    session["selected_set"] = set_number

    session["stage"] = "ready"

    session.pop(
        "test_start_time",
        None
    )


    return render_template(

        "set_ready.html",

        set_number=set_number,

        duration=TEST_DURATION

    )


# ============================================================
# BEGIN TEST
# ============================================================

@app.route("/begin", methods=["POST"])
def begin():

    if (
        session.get("stage") != "ready"
        or session.get("selected_set") not in SETS
    ):

        return redirect("/sets")


    session["test_start_time"] = time.time()

    session["stage"] = "test"

    session["test_active"] = True

    session["submitted"] = False


    return redirect(
        "/test"
    )


# ============================================================
# TEST PAGE
# ============================================================

@app.route("/test")
def test():

    if (
        session.get("stage") != "test"
        or not session.get("test_active")
    ):

        return redirect("/sets")


    elapsed = current_elapsed()


    remaining = max(
        0,
        TEST_DURATION - elapsed
    )


    # --------------------------------------------------------
    # Time expired
    # --------------------------------------------------------

    if remaining <= 0:

        return _finalize(

            score=0,

            status="Time Expired",

            details=[]

        )


    selected = SETS[
        session["selected_set"]
    ]


    # --------------------------------------------------------
    # Send questions to test.html
    # --------------------------------------------------------

    return render_template(

        "test.html",

        questions=selected["questions"],

        remaining=remaining,

        name=session["name"],

        team_name=session["team_name"],

        set_number=session["selected_set"]

    )


# ============================================================
# TERMINATE TEST
# ============================================================

@app.route("/terminate", methods=["POST"])
def terminate():

    if (
        session.get("stage") != "test"
        or not session.get("test_active")
    ):

        return jsonify(
            {
                "success": False
            }
        )


    _finalize(

        score=0,

        status="Eliminated - Anti-Cheat Violation",

        details=[],

        redirect_after=False

    )


    return jsonify(

        {
            "success": True,

            "url": "/thank-you"
        }

    )


# ============================================================
# FINALIZE RESULT
# ============================================================

def _finalize(
    score,
    status,
    details,
    answers=None,
    redirect_after=True
):

    if not session.get(
        "test_active"
    ):

        if redirect_after:

            return redirect(
                "/thank-you"
            )

        return None


    elapsed = min(
        current_elapsed(),
        TEST_DURATION
    )


    set_number = session.get(
        "selected_set"
    )


    max_score = (

        len(
            SETS[set_number]["questions"]
        )

        * MAX_SCORE_PER_QUESTION

    )


    # --------------------------------------------------------
    # Save result
    # --------------------------------------------------------

    save_result(

        session["name"],

        session["email"],

        session["team_name"],

        set_number,

        score,

        max_score,

        elapsed,

        status,

        answers or {}

    )


    # --------------------------------------------------------
    # Store result in session
    # --------------------------------------------------------

    session["test_active"] = False

    session["stage"] = "finished"

    session["score"] = score

    session["time_taken"] = elapsed

    session["status"] = status

    session["details"] = details


    if redirect_after:

        return redirect(
            "/thank-you"
        )


    return None


# ============================================================
# SUBMIT TEST
# ============================================================

@app.route("/submit", methods=["POST"])
def submit_test():

    if (
        session.get("stage") != "test"
        or not session.get("test_active")
    ):

        return redirect(
            "/thank-you"
        )


    # --------------------------------------------------------
    # Get selected set and preserve submitted answers
    # --------------------------------------------------------

    selected = SETS[
        session["selected_set"]
    ]

    submitted_answers = {
        str(question["id"]): request.form.get(
            f"answer{question['id']}",
            ""
        )
        for question in selected["questions"]
    }

    # --------------------------------------------------------
    # Check timer
    # --------------------------------------------------------

    elapsed = current_elapsed()


    if elapsed >= TEST_DURATION:

        return _finalize(

            score=0,

            status="Time Expired",

            details=[],

            answers=submitted_answers

        )


    total = 0

    details = []


    # --------------------------------------------------------
    # Evaluate every question
    # --------------------------------------------------------

    for question in selected["questions"]:

        answer = submitted_answers[str(question["id"])]


        score, passed, message = evaluate_question(

            question,

            answer

        )


        total += score


        details.append(

            {

                "id": question["id"],

                "title": question["title"],

                "score": score,

                "passed": passed,

                "message": message

            }

        )


    # --------------------------------------------------------
    # Save final result
    # --------------------------------------------------------

    return _finalize(

        score=total,

        status="Completed",

        details=details,

        answers=submitted_answers

    )


# ============================================================
# THANK YOU / RESULT
# ============================================================

@app.route("/thank-you")
def thank_you():

    if not session.get(
        "name"
    ):

        return redirect("/login")


    status = session.get("status", "Completed")
    eliminated = status != "Completed"

    return render_template(

        "result.html",

        name=session.get("name", "Participant"),

        team_name=session.get("team_name", ""),

        heading="You Have Been Eliminated" if eliminated else "Thank You for Participating!",

        page_title="Eliminated" if eliminated else "Thank You",

        eliminated=eliminated

    )


# ============================================================
# LEADERBOARD
# ============================================================

@app.route("/leaderboard")
def leaderboard():

    return redirect(
        "/admin"
    )


# ============================================================
# ADMIN
# ============================================================

@app.route("/admin", methods=["GET", "POST"])
def admin():

    if request.method == "POST":

        password = request.form.get(
            "password",
            ""
        )


        if password == ADMIN_PASSWORD:

            session[
                "admin_authenticated"
            ] = True

        else:

            return render_template(

                "admin_login.html",

                error=(
                    "Incorrect admin password."
                )

            )


    if not session.get(
        "admin_authenticated"
    ):

        return render_template(
            "admin_login.html"
        )


    return render_template(

        "leaderboard.html",

        results=get_leaderboard(),

        public=False,

        dustbin_error=request.args.get("error") == "dustbin",
        show_error=request.args.get("error") == "show"

    )


# ============================================================
# DELETE ONE RESULT (ADMIN + DUSTBIN PASSWORD)
# ============================================================

@app.route("/admin/delete/<int:result_id>", methods=["POST"])
def admin_delete_result(result_id):

    if not admin_required():
        return redirect("/admin")

    password = request.form.get("dustbin_password", "")

    if password != DUSTBIN_PASSWORD:
        return redirect("/admin?error=dustbin")

    delete_result(result_id)
    return redirect("/admin")


# ============================================================
# CLEAR ALL RESULTS (ADMIN + DUSTBIN PASSWORD)
# ============================================================

@app.route("/admin/clear", methods=["POST"])
def admin_clear_results():

    if not admin_required():
        return redirect("/admin")

    password = request.form.get("dustbin_password", "")

    if password != DUSTBIN_PASSWORD:
        return redirect("/admin?error=dustbin")

    clear_results()
    return redirect("/admin")


# ============================================================
# SHOW ONE PARTICIPANT'S ANSWERS
# ============================================================

@app.route("/admin/show/<int:result_id>", methods=["POST"])
def admin_show_answers(result_id):

    if not admin_required():
        return redirect("/admin")

    password = request.form.get("show_password", "")
    if password != SHOW_ANSWERS_PASSWORD:
        return redirect("/admin?error=show")

    row = get_result(result_id)
    if row is None:
        return redirect("/admin")

    try:
        answers = json.loads(row["answers_json"] or "{}")
    except (TypeError, json.JSONDecodeError):
        answers = {}

    question_map = {}
    for set_data in SETS.values():
        for question in set_data["questions"]:
            question_map[str(question["id"])] = question

    answer_rows = []
    for question_id, answer in answers.items():
        question = question_map.get(str(question_id))
        answer_rows.append({
            "id": question_id,
            "title": question["title"] if question else f"Question {question_id}",
            "answer": answer or "(No answer submitted)",
        })

    answer_rows.sort(key=lambda item: int(item["id"]) if str(item["id"]).isdigit() else str(item["id"]))

    return render_template(
        "admin_answers.html",
        result=row,
        answers=answer_rows,
        show_password_used=True
    )


# ============================================================
# DOWNLOAD ONE PARTICIPANT'S ANSWERS
# ============================================================

@app.route("/admin/download/<int:result_id>.xlsx")
def download_answers(result_id):

    if not admin_required():
        return redirect("/admin")

    row = get_result(result_id)
    if row is None:
        return redirect("/admin")

    try:
        answers = json.loads(row["answers_json"] or "{}")
    except (TypeError, json.JSONDecodeError):
        answers = {}

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    question_map = {}
    for set_data in SETS.values():
        for question in set_data["questions"]:
            question_map[str(question["id"])] = question

    wb = Workbook()
    ws = wb.active
    ws.title = "Participant Answers"

    ws.append(["Team Name", row["team_name"]])
    ws.append(["Participant", row["student_name"]])
    ws.append(["Email", row["email"]])
    ws.append(["Set", f"Set {row['set_number']}"])
    ws.append(["Score", f"{row['score']}/{row['max_score']}"])
    ws.append([])
    ws.append(["Question", "Title", "Submitted Answer"])
    for cell in ws[7]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for qid in sorted(answers, key=lambda x: int(x) if str(x).isdigit() else str(x)):
        question = question_map.get(str(qid))
        ws.append([
            f"Q{qid}",
            question["title"] if question else f"Question {qid}",
            answers.get(qid, "")
        ])

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 90
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    safe_team = re.sub(r"[^A-Za-z0-9_-]+", "_", row["team_name"]).strip("_") or "team"
    filename = f"{safe_team}_Set_{row['set_number']}_Answers.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ============================================================
# ADMIN LOGOUT
# ============================================================

@app.route("/admin/logout")
def admin_logout():

    # End the admin session completely and return to the participant
    # username/details page as requested.
    session.clear()

    return redirect(
        "/login"
    )


# ============================================================
# EXPORT EXCEL
# ============================================================

@app.route("/admin/export.xlsx")
def export_excel():

    if not session.get(
        "admin_authenticated"
    ):

        return redirect(
            "/admin"
        )


    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment


    rows = get_leaderboard()


    wb = Workbook()


    ws = wb.active

    ws.title = "Result Sheet"


    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [

        "Rank",

        "Team Name",

        "Participant",

        "Email",

        "Set",

        "Score",

        "Max Marks",

        "Time",

        "Status",

        "Submitted At"

    ]


    ws.append(
        headers
    )


    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


    # --------------------------------------------------------
    # Add results
    # --------------------------------------------------------

    for rank, row in enumerate(
        rows,
        1
    ):

        minutes, seconds = divmod(

            row["time_taken"],

            60

        )


        ws.append(

            [

                rank,

                row["team_name"],

                row["student_name"],

                row["email"],

                f"Set {row['set_number']}",

                row["score"],

                row["max_score"],

                f"{minutes:02d}:{seconds:02d}",

                row["status"],

                row["submitted_at"]

            ]

        )


    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = [

        8,

        20,

        24,

        30,

        10,

        10,

        12,

        12,

        25,

        22

    ]


    for index, width in enumerate(
        widths,
        1
    ):

        ws.column_dimensions[
            chr(64 + index)
        ].width = width


    # --------------------------------------------------------
    # Create Excel file
    # --------------------------------------------------------

    output = BytesIO()


    wb.save(
        output
    )


    output.seek(0)


    return send_file(

        output,

        as_attachment=True,

        download_name=(
            "cyber_debug_result_sheet.xlsx"
        ),

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    )


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    return jsonify(
        {
            "status": "ok"
        }
    )


# ============================================================
# START FLASK
# ============================================================

if __name__ == "__main__":

    initialize_database()

    app.run(

        debug=True,

        host="127.0.0.1",

        port=5000

    )